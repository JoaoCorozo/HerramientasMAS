"""Consolidar cambios cliente EMIN: nómina antigua + nueva → unión con parche."""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import BinaryIO

import openpyxl
import pandas as pd
from openpyxl.styles import PatternFill

from emin_comparador import (
    CAMPOS,
    _estilo_cuerpo,
    _estilo_header,
    _leer_archivo,
    _normalizar_dataframe,
)

CAMPOS_IDENTIDAD = ("RUT", "NOMBRE", "APELLIDO")

FILL_CONSOLIDADA = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
FILL_ACTUALIZADOS = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
FILL_NUEVOS = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")


def _fila_campos(row) -> dict[str, str]:
    return {c: str(row.get(c, "") or "") for c in CAMPOS}


def _merge_actualizado(row_antigua, row_nueva) -> dict[str, str]:
    """Identidad desde antigua; resto de columnas desde nueva."""
    out = _fila_campos(row_nueva)
    for campo in CAMPOS_IDENTIDAD:
        out[campo] = str(row_antigua.get(campo, "") or "")
    return out


def consolidar_nominas(
    ruta_antigua: Path | str,
    ruta_nueva: Path | str,
) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]], dict]:
    df_ant = _normalizar_dataframe(_leer_archivo(ruta_antigua), "antigua")
    df_nue = _normalizar_dataframe(_leer_archivo(ruta_nueva), "nueva")

    idx_nueva = df_nue.set_index("RUT_norm", drop=False)
    ruts_antigua = set(df_ant["RUT_norm"])
    ruts_nueva = set(df_nue["RUT_norm"])

    consolidada: list[dict[str, str]] = []
    actualizados: list[dict[str, str]] = []
    nuevos: list[dict[str, str]] = []

    for _, row in df_ant.iterrows():
        rut = row["RUT_norm"]
        if rut in ruts_nueva:
            row_n = idx_nueva.loc[rut]
            if isinstance(row_n, pd.DataFrame):
                row_n = row_n.iloc[0]
            merged = _merge_actualizado(row, row_n)
            consolidada.append(merged)
            actualizados.append(merged)
        else:
            consolidada.append(_fila_campos(row))

    for _, row in df_nue.iterrows():
        rut = row["RUT_norm"]
        if rut in ruts_antigua:
            continue
        fila = _fila_campos(row)
        consolidada.append(fila)
        nuevos.append(fila)

    stats = {
        "antigua": int(len(df_ant)),
        "nueva": int(len(df_nue)),
        "consolidada": len(consolidada),
        "actualizados": len(actualizados),
        "nuevos": len(nuevos),
        "solo_antigua": int(len(ruts_antigua - ruts_nueva)),
    }
    return consolidada, actualizados, nuevos, stats


def _escribir_hoja(wb, titulo: str, filas: list[dict[str, str]], fill: PatternFill):
    ws = wb.create_sheet(titulo)
    ws.append(list(CAMPOS))
    _estilo_header(ws, fill)
    for fila in filas:
        ws.append([fila.get(c, "") for c in CAMPOS])
    _estilo_cuerpo(ws)


def procesar_consolidacion(
    ruta_antigua: Path | str,
    ruta_nueva: Path | str,
    destino: Path | str | BinaryIO,
) -> dict:
    consolidada, actualizados, nuevos, stats = consolidar_nominas(ruta_antigua, ruta_nueva)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _escribir_hoja(wb, "Consolidada", consolidada, FILL_CONSOLIDADA)
    _escribir_hoja(wb, "Actualizados", actualizados, FILL_ACTUALIZADOS)
    _escribir_hoja(wb, "Nuevos", nuevos, FILL_NUEVOS)
    wb.save(destino)
    return stats


def generar_reporte_bytes(
    ruta_antigua: Path | str,
    ruta_nueva: Path | str,
) -> tuple[bytes, str, dict]:
    buf = io.BytesIO()
    stats = procesar_consolidacion(ruta_antigua, ruta_nueva, buf)
    filename = f"EMIN_Consolidar_cambios_cliente_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return buf.getvalue(), filename, stats
