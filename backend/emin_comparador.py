"""Comparador de nóminas EMIN: cliente vs corregida → Corregidos / Consultar a Cliente / Mapeo."""

from __future__ import annotations

import csv
import io
import re
import warnings
from datetime import datetime
from pathlib import Path
from typing import BinaryIO

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CAMPOS = [
    "RUT",
    "NOMBRE",
    "APELLIDO",
    "CORREO",
    "GENERO",
    "FECHA DE INGRESO",
    "FAMILIA CARGO",
    "CARGO",
    "GERENCIA",
    "SUBGERENCIA",
    "AREA",
    "CENTRO DE COSTO",
    "EMPRESA",
    "RUT PRIMER JEFE",
    "NOMBRE PRIMER JEFE",
    "RUT SEGUNDO JEFE",
    "NOMBRE SEGUNDO JEFE",
    "TIPO DE EVALUACIÓN",
    "¿PARTICIPA EVALUACIÓN POTENCIAL?",
]

# Alias de encabezados frecuentes → nombre canónico
ALIAS_CAMPOS: dict[str, list[str]] = {
    "RUT": ["RUT", "Rut", "rut", "RUN"],
    "NOMBRE": ["NOMBRE", "NOMBRES", "Nombre", "Nombres"],
    "APELLIDO": ["APELLIDO", "APELLIDOS", "Apellido", "Apellidos"],
    "CORREO": ["CORREO", "EMAIL", "Email", "Correo"],
    "GENERO": ["GENERO", "GÉNERO", "Genero", "Sexo"],
    "FECHA DE INGRESO": ["FECHA DE INGRESO", "Fecha de Ingreso", "FECHA INGRESO"],
    "FAMILIA CARGO": ["FAMILIA CARGO", "Familia Cargo", "FAMILIA DE CARGO"],
    "CARGO": ["CARGO", "Cargo", "NOMBRE DE CARGO"],
    "GERENCIA": ["GERENCIA", "Gerencia"],
    "SUBGERENCIA": ["SUBGERENCIA", "Subgerencia"],
    "AREA": ["AREA", "ÁREA", "Area", "Área"],
    "CENTRO DE COSTO": ["CENTRO DE COSTO", "CENTRO DE COSTOS", "Centro de Costo", "Centro Costo"],
    "EMPRESA": ["EMPRESA", "Empresa"],
    "RUT PRIMER JEFE": [
        "RUT PRIMER JEFE",
        "RUT 1ER JEFE",
        "Rut Primer Jefe",
        "RUT primer jefe (opcional)",
    ],
    "NOMBRE PRIMER JEFE": [
        "NOMBRE PRIMER JEFE",
        "NOMBRE 1ER JEFE",
        "Nombre Primer Jefe",
        "Nombre primer jefe (opcional)",
        "NOMBRE primer jefe (opcional)",
    ],
    "RUT SEGUNDO JEFE": [
        "RUT SEGUNDO JEFE",
        "RUT 2DO JEFE",
        "Rut Segundo Jefe",
        "RUT segundo jefe (opcional)",
    ],
    "NOMBRE SEGUNDO JEFE": [
        "NOMBRE SEGUNDO JEFE",
        "NOMBRE 2DO JEFE",
        "Nombre Segundo Jefe",
        "Nombre segundo jefe (opcional)",
    ],
    "TIPO DE EVALUACIÓN": ["TIPO DE EVALUACIÓN", "TIPO DE EVALUACION", "Tipo de Evaluación"],
    "¿PARTICIPA EVALUACIÓN POTENCIAL?": [
        "¿PARTICIPA EVALUACIÓN POTENCIAL?",
        "PARTICIPA EVALUACIÓN POTENCIAL",
        "PARTICIPA EVALUACION POTENCIAL",
        "¿PARTICIPA EVALUACION POTENCIAL?",
    ],
}

MARCADORES_NO_EXISTE = frozenset(
    {
        "999 no existe",
        "9999 no existe",
        "999 noexiste",
        "9999 noexiste",
    }
)

# Campos siempre obligatorios (RUT PRIMER JEFE tiene regla condicional aparte)
CAMPOS_OBLIGATORIOS = [
    "RUT",
    "NOMBRE",
    "APELLIDO",
    "CORREO",
    "GENERO",
    "FECHA DE INGRESO",
    "CARGO",
    "FAMILIA CARGO",
    "AREA",
    "CENTRO DE COSTO",
    "EMPRESA",
    "TIPO DE EVALUACIÓN",
    "¿PARTICIPA EVALUACIÓN POTENCIAL?",
]

# Nombre visible en ERROR (alineado al plan / Excel del cliente)
NOMBRE_ERROR_CAMPO = {
    "FAMILIA CARGO": "FAMILIA DE CARGO",
    "TIPO DE EVALUACIÓN": "TIPO DE EVALUACION",
}

FILL_HEADER_OK = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
FILL_HEADER_ASK = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
FILL_HEADER_MAP = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
FILL_ERROR_CELL = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")


def normalizar_rut(valor) -> str:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    s = str(valor).strip()
    if s.endswith(".0"):
        s = s[:-2]
    s = s.replace("\xa0", "").replace(" ", "")
    return re.sub(r"[^0-9Kk]", "", s).upper()


def _es_no_existe(valor) -> bool:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return False
    s = " ".join(str(valor).strip().casefold().split())
    s_compact = s.replace(" ", "")
    return s in MARCADORES_NO_EXISTE or s_compact in MARCADORES_NO_EXISTE


def _parsear_fecha_ingreso(valor: str) -> datetime | None:
    """Parsea fecha; None si vacía o calendario imposible (ej. 30/02/2012)."""
    s = str(valor or "").strip()
    if not s or s.casefold() in ("nan", "none", "nat"):
        return None
    s0 = s.split()[0]
    if "T" in s0:
        s0 = s0.split("T")[0]

    for fmt in (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d/%m/%y",
        "%d-%m-%y",
        "%Y/%m/%d",
        "%d.%m.%Y",
    ):
        try:
            return datetime.strptime(s0, fmt)
        except ValueError:
            continue

    # Serial Excel (días desde 1899-12-30)
    try:
        n = float(s.replace(",", "."))
        if 1 < n < 100000:
            from datetime import timedelta

            return datetime(1899, 12, 30) + timedelta(days=n)
    except ValueError:
        pass
    return None


def _es_fecha_ingreso_invalida(valor: str) -> bool:
    """True si hay valor pero no es fecha real o el año es < 1980."""
    s = str(valor or "").strip()
    if not s or s.casefold() in ("nan", "none", "nat"):
        return False
    dt = _parsear_fecha_ingreso(s)
    if dt is None:
        return True
    return dt.year < 1980


def _detectar_encoding(ruta: Path | str) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            with open(ruta, encoding=encoding) as f:
                f.readline()
            return encoding
        except UnicodeDecodeError:
            continue
    return "latin-1"


def _detectar_separador(texto: str) -> str:
    primera = texto.splitlines()[0] if texto else ""
    try:
        dialecto = csv.Sniffer().sniff(texto[:8192], delimiters=";,")
        if dialecto.delimiter in (";", ","):
            return dialecto.delimiter
    except csv.Error:
        pass
    return ";" if primera.count(";") >= primera.count(",") else ","


def _leer_csv(ruta: Path | str) -> pd.DataFrame:
    encoding = _detectar_encoding(ruta)
    with open(ruta, encoding=encoding, errors="replace") as f:
        texto = f.read()
    sep = _detectar_separador(texto)
    kwargs = dict(sep=sep, dtype=str, engine="python", quotechar='"', skipinitialspace=True)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", category=pd.errors.ParserWarning)
        try:
            return pd.read_csv(io.StringIO(texto), on_bad_lines="warn", **kwargs)
        except TypeError:
            return pd.read_csv(
                io.StringIO(texto),
                error_bad_lines=False,
                warn_bad_lines=False,
                **kwargs,
            )


def _leer_archivo(ruta: Path | str) -> pd.DataFrame:
    path = Path(ruta)
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xls"):
        return pd.read_excel(path, dtype=str)
    if ext == ".csv":
        return _leer_csv(path)
    raise ValueError("Solo se permiten archivos Excel (.xlsx/.xls) o CSV (.csv).")


def _clave_columna(nombre: str) -> str:
    """Normaliza nombre de columna para match (ignora (opcional), signos, etc.)."""
    s = str(nombre).casefold()
    s = re.sub(r"\(.*?\)", "", s)  # quita (opcional) y similares
    s = s.replace("opcional", "")
    return re.sub(r"[^a-z0-9]", "", s)


def _buscar_columna(df: pd.DataFrame, candidatos: list[str]) -> str | None:
    mapa = {str(c).strip().casefold(): c for c in df.columns}
    for cand in candidatos:
        key = cand.strip().casefold()
        if key in mapa:
            return mapa[key]
    for col in df.columns:
        limpio = _clave_columna(col)
        for cand in candidatos:
            if limpio and limpio == _clave_columna(cand):
                return col
    return None


def _normalizar_dataframe(
    df: pd.DataFrame,
    etiqueta: str,
    *,
    dedupe: bool = True,
    drop_empty_rut: bool = True,
) -> pd.DataFrame:
    col_rut = _buscar_columna(df, ALIAS_CAMPOS["RUT"])
    if not col_rut:
        raise ValueError(f"No se encontró la columna RUT en la nómina {etiqueta}.")

    out = pd.DataFrame()
    for campo in CAMPOS:
        src = _buscar_columna(df, ALIAS_CAMPOS.get(campo, [campo]))
        out[campo] = df[src].fillna("").astype(str).str.strip() if src else ""
        out[campo] = out[campo].replace({"nan": "", "None": ""})

    # Fila Excel (fila 1 = encabezado → primera dato = 2), para cruzar con el analizador
    out["_fila_excel"] = list(range(2, len(out) + 2))
    out["RUT_norm"] = out["RUT"].apply(normalizar_rut)
    if drop_empty_rut:
        out = out[out["RUT_norm"] != ""].copy()
    if dedupe:
        out = out.drop_duplicates(subset=["RUT_norm"], keep="first")
    return out


def _clave_mapeo(columna: str, valor_antes: str) -> str:
    return f"{columna}||{valor_antes}"


def _nombre_error_campo(campo: str) -> str:
    return NOMBRE_ERROR_CAMPO.get(campo, campo)


def _es_valor_no(valor) -> bool:
    """True si el valor es No / NO / no (sin importar mayúsculas)."""
    return str(valor or "").strip().casefold() == "no"


def _permite_rut_primer_jefe_vacio(row_cli) -> bool:
    """RUT PRIMER JEFE puede ir vacío solo si TIPO y PARTICIPA son ambos 'No'."""
    tipo = row_cli.get("TIPO DE EVALUACIÓN", "")
    participa = row_cli.get("¿PARTICIPA EVALUACIÓN POTENCIAL?", "")
    return _es_valor_no(tipo) and _es_valor_no(participa)


def _canonizar_columna(nombre: str) -> str | None:
    """Mapea nombre de columna del analizador al canónico de CAMPOS."""
    if not nombre:
        return None
    clave = _clave_columna(nombre)
    for campo, aliases in ALIAS_CAMPOS.items():
        for alias in aliases + [campo]:
            if _clave_columna(alias) == clave:
                return campo
    for campo in CAMPOS:
        if _clave_columna(campo) == clave:
            return campo
    return None


def _validar_fila_cliente(
    row_cli, ruts_duplicados: set[str]
) -> tuple[list[dict[str, str | None]], list[str]]:
    """Devuelve (errores estructurados {columna, mensaje}, celdas_rojas)."""
    errores: list[dict[str, str | None]] = []
    celdas_rojas: list[str] = []

    rut_norm = str(row_cli.get("RUT_norm", "") or "")
    if rut_norm and rut_norm in ruts_duplicados:
        errores.append({"columna": None, "mensaje": "usuario duplicado"})

    for campo in CAMPOS_OBLIGATORIOS:
        val = str(row_cli.get(campo, "") or "").strip()
        if not val:
            errores.append(
                {
                    "columna": campo,
                    "mensaje": f"Campo {_nombre_error_campo(campo)} sin informacion",
                }
            )
            celdas_rojas.append(campo)

    rut_jefe = str(row_cli.get("RUT PRIMER JEFE", "") or "").strip()
    if not rut_jefe and not _permite_rut_primer_jefe_vacio(row_cli):
        errores.append({"columna": "RUT PRIMER JEFE", "mensaje": "Campo RUT PRIMER JEFE sin informacion"})
        celdas_rojas.append("RUT PRIMER JEFE")

    fecha_val = str(row_cli.get("FECHA DE INGRESO", "") or "").strip()
    if fecha_val and _es_fecha_ingreso_invalida(fecha_val):
        errores.append({"columna": "FECHA DE INGRESO", "mensaje": "Fecha de ingreso no valida"})
        if "FECHA DE INGRESO" not in celdas_rojas:
            celdas_rojas.append("FECHA DE INGRESO")

    return errores, celdas_rojas


def _leer_analizador(ruta: Path | str) -> list[dict]:
    """Lee Fila / Columna / Error del Excel del analizador de la plataforma."""
    df = _leer_archivo(ruta)
    col_fila = _buscar_columna(df, ["Fila", "FILA", "fila", "Row", "ROW"])
    col_columna = _buscar_columna(df, ["Columna", "COLUMNA", "columna", "Campo", "CAMPO"])
    col_error = _buscar_columna(df, ["Error", "ERROR", "error", "Mensaje", "MENSAJE"])
    if not col_fila or not col_columna or not col_error:
        raise ValueError(
            "El archivo del analizador debe tener columnas Fila, Columna y Error."
        )

    items: list[dict] = []
    for _, row in df.iterrows():
        raw_fila = str(row.get(col_fila, "") or "").strip()
        if not raw_fila or raw_fila.casefold() in ("nan", "none"):
            continue
        try:
            fila = int(float(raw_fila.replace(",", ".")))
        except ValueError:
            continue
        columna_raw = str(row.get(col_columna, "") or "").strip()
        error = str(row.get(col_error, "") or "").strip()
        if not columna_raw or not error:
            continue
        columna = _canonizar_columna(columna_raw) or columna_raw
        items.append(
            {
                "fila": fila,
                "columna": columna,
                "columna_raw": columna_raw,
                "error": error,
                "id": f"{fila}|{columna}",
            }
        )
    return items


def _mensajes_error(errores: list[dict[str, str | None]]) -> str:
    return "; ".join(str(e["mensaje"]) for e in errores if e.get("mensaje"))


def _estilo_header(ws, fill: PatternFill):
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = font_header
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30


def _estilo_cuerpo(ws):
    font_body = Font(name="Segoe UI", size=10)
    border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = font_body
            cell.border = border
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col[: min(80, len(col))])
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 3, 12), 45)
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _construir_filas(
    ruta_cliente: Path | str,
    ruta_corregida: Path | str,
) -> tuple[list[dict], dict[str, dict[str, str]], pd.DataFrame, pd.DataFrame]:
    """Arma filas internas del comparador (antes de analizador / Excel)."""
    df_cli = _normalizar_dataframe(
        _leer_archivo(ruta_cliente),
        "cliente",
        dedupe=False,
        drop_empty_rut=False,
    )
    df_corr = _normalizar_dataframe(_leer_archivo(ruta_corregida), "corregida")

    rut_counts = df_cli.loc[df_cli["RUT_norm"] != "", "RUT_norm"].value_counts()
    ruts_duplicados = set(rut_counts[rut_counts > 1].index)

    corr_idx = df_corr.set_index("RUT_norm", drop=False)
    ruts_corr_procesados: set[str] = set()
    mapeo: dict[str, dict[str, str]] = {}
    filas: list[dict] = []

    def _aplicar_no_existe(row_cli, row_corr, fila: dict) -> None:
        for campo in CAMPOS:
            val_corr = str(row_corr.get(campo, "") or "").strip()
            val_cli = str(row_cli.get(campo, "") or "").strip() if row_cli is not None else ""
            if _es_no_existe(val_corr):
                fila["errores"].append({"columna": campo, "mensaje": f"No existe {campo}"})
                fila["campos"][campo] = val_cli
                if campo not in fila["celdas_rojas"]:
                    fila["celdas_rojas"].append(campo)
                clave = _clave_mapeo(campo, val_cli)
                fila["claves_info"].append((campo, clave))
                if clave not in mapeo:
                    mapeo[clave] = {
                        "Columna": campo,
                        "Valor antes": val_cli,
                        "RUT ejemplo": str(row_corr.get("RUT", "") or ""),
                    }
            else:
                fila["campos"][campo] = val_corr

    for _, row_cli in df_cli.iterrows():
        rut_norm = str(row_cli.get("RUT_norm", "") or "")
        errores_cli, celdas_rojas = _validar_fila_cliente(row_cli, ruts_duplicados)
        fila: dict = {
            "_fila_excel": int(row_cli.get("_fila_excel") or 0),
            "_rut_norm": rut_norm,
            "_origen": "cliente",
            "campos": {},
            "errores": list(errores_cli),
            "celdas_rojas": list(celdas_rojas),
            "claves_info": [],
            "errores_cli": bool(errores_cli),
        }

        row_corr = None
        if rut_norm and rut_norm in corr_idx.index:
            row_corr = corr_idx.loc[rut_norm]
            if isinstance(row_corr, pd.DataFrame):
                row_corr = row_corr.iloc[0]
            ruts_corr_procesados.add(rut_norm)

        if row_corr is not None:
            _aplicar_no_existe(row_cli, row_corr, fila)
        else:
            for campo in CAMPOS:
                fila["campos"][campo] = str(row_cli.get(campo, "") or "").strip()

        if fila["errores"]:
            if errores_cli:
                for campo in CAMPOS:
                    fila["campos"][campo] = str(row_cli.get(campo, "") or "").strip()
            fila["destino"] = "consultar"
            filas.append(fila)
        elif row_corr is not None:
            fila["destino"] = "corregidos"
            filas.append(fila)

    for rut_norm, row_corr in corr_idx.iterrows():
        if rut_norm in ruts_corr_procesados:
            continue
        fila = {
            "_fila_excel": 0,
            "_rut_norm": rut_norm,
            "_origen": "corregida",
            "campos": {},
            "errores": [],
            "celdas_rojas": [],
            "claves_info": [],
            "errores_cli": False,
        }
        for campo in CAMPOS:
            val_corr = str(row_corr.get(campo, "") or "").strip()
            if _es_no_existe(val_corr):
                fila["errores"].append({"columna": campo, "mensaje": f"No existe {campo}"})
                fila["campos"][campo] = ""
                fila["celdas_rojas"].append(campo)
                clave = _clave_mapeo(campo, "")
                fila["claves_info"].append((campo, clave))
                if clave not in mapeo:
                    mapeo[clave] = {
                        "Columna": campo,
                        "Valor antes": "",
                        "RUT ejemplo": str(row_corr.get("RUT", "") or ""),
                    }
            else:
                fila["campos"][campo] = val_corr
        fila["destino"] = "consultar" if fila["errores"] else "corregidos"
        filas.append(fila)

    return filas, mapeo, df_cli, df_corr


def _aplicar_analizador(
    filas: list[dict],
    df_cli: pd.DataFrame,
    df_corr: pd.DataFrame,
    items_analizador: list[dict],
    decisiones: dict[str, str] | None,
) -> list[dict]:
    """
    Cruza errores del analizador.
    - Misma columna ya en error del comparador → reemplaza mensaje por el del analizador.
    - Columna sin error del comparador → pending (o aplica decisión).
    """
    decisiones = decisiones or {}
    pending: list[dict] = []
    corr_idx = df_corr.set_index("RUT_norm", drop=False) if len(df_corr) else df_corr

    # Indexar filas cliente por _fila_excel
    por_fila: dict[int, list[dict]] = {}
    for f in filas:
        if f.get("_origen") == "cliente" and f.get("_fila_excel"):
            por_fila.setdefault(int(f["_fila_excel"]), []).append(f)

    # También indexar df_cli por fila excel (por si la fila no entró al resultado)
    cli_por_fila = {
        int(r["_fila_excel"]): r for _, r in df_cli.iterrows() if r.get("_fila_excel")
    }

    for item in items_analizador:
        fila_n = int(item["fila"])
        columna = item["columna"]
        error_txt = item["error"]
        item_id = item["id"]

        candidatas = por_fila.get(fila_n, [])
        row_cli = cli_por_fila.get(fila_n)

        if not candidatas and row_cli is None:
            continue

        # Usar la primera candidata o crear desde cliente
        if candidatas:
            fila = candidatas[0]
        else:
            # Fila cliente sin destino previo: crear shell
            rut_norm = str(row_cli.get("RUT_norm", "") or "")
            campos = {c: str(row_cli.get(c, "") or "").strip() for c in CAMPOS}
            fila = {
                "_fila_excel": fila_n,
                "_rut_norm": rut_norm,
                "_origen": "cliente",
                "campos": campos,
                "errores": [],
                "celdas_rojas": [],
                "claves_info": [],
                "errores_cli": False,
                "destino": None,
            }
            filas.append(fila)
            por_fila.setdefault(fila_n, []).append(fila)

        cols_con_error = {e.get("columna") for e in fila["errores"] if e.get("columna")}
        if columna in cols_con_error:
            for e in fila["errores"]:
                if e.get("columna") == columna:
                    e["mensaje"] = error_txt
            if columna not in fila["celdas_rojas"]:
                fila["celdas_rojas"].append(columna)
            fila["destino"] = "consultar"
            continue

        # No hay error del comparador en esa columna → decisión manual
        destino = decisiones.get(item_id)
        if destino not in ("consultar", "corregidos"):
            if row_cli is not None:
                rut_cli = str(row_cli.get("RUT", "") or "")
                nom_cli = str(row_cli.get("NOMBRE", "") or "")
                ape_cli = str(row_cli.get("APELLIDO", "") or "")
                val_cli = str(row_cli.get(columna, "") or "") if columna in CAMPOS else ""
                rut_norm = str(row_cli.get("RUT_norm", "") or "")
            else:
                rut_cli = str(fila["campos"].get("RUT", ""))
                nom_cli = str(fila["campos"].get("NOMBRE", ""))
                ape_cli = str(fila["campos"].get("APELLIDO", ""))
                val_cli = str(fila["campos"].get(columna, ""))
                rut_norm = str(fila.get("_rut_norm", ""))

            val_corr = ""
            rut_corr = nom_corr = ape_corr = ""
            row_corr = None
            if rut_norm and len(df_corr) and rut_norm in corr_idx.index:
                row_corr = corr_idx.loc[rut_norm]
                if isinstance(row_corr, pd.DataFrame):
                    row_corr = row_corr.iloc[0]
                rut_corr = str(row_corr.get("RUT", "") or "")
                nom_corr = str(row_corr.get("NOMBRE", "") or "")
                ape_corr = str(row_corr.get("APELLIDO", "") or "")
                val_corr = str(row_corr.get(columna, "") or "") if columna in CAMPOS else ""

            mismo_valor = val_cli.strip() == val_corr.strip()
            if row_corr is None:
                porque = (
                    f'La fila {fila_n} corresponde al rut "{rut_cli}" en la nómina cliente. '
                    f"No hay coincidencia por RUT en la nómina corregida, pero el analizador indica "
                    f'en {columna}: "{error_txt}".'
                )
            elif mismo_valor:
                porque = (
                    f'La fila {fila_n} corresponde al rut "{rut_cli}" en la nómina cliente. '
                    f"El valor de {columna} no presenta cambios en la nómina corregida, "
                    f'pero el analizador indica: "{error_txt}".'
                )
            else:
                porque = (
                    f'La fila {fila_n} corresponde al rut "{rut_cli}" en la nómina cliente. '
                    f"El comparador no marcó error en {columna}, "
                    f'pero el analizador indica: "{error_txt}".'
                )

            pending.append(
                {
                    "id": item_id,
                    "fila": fila_n,
                    "columna": columna,
                    "error_analizador": error_txt,
                    "motivo": porque,
                    "cliente": {
                        "rut": rut_cli,
                        "nombre": nom_cli,
                        "apellido": ape_cli,
                        "valor": val_cli,
                    },
                    "corregida": {
                        "rut": rut_corr,
                        "nombre": nom_corr,
                        "apellido": ape_corr,
                        "valor": val_corr,
                    },
                }
            )
            continue

        if destino == "consultar":
            fila["errores"].append({"columna": columna, "mensaje": error_txt})
            if columna not in fila["celdas_rojas"]:
                fila["celdas_rojas"].append(columna)
            if row_cli is not None:
                for campo in CAMPOS:
                    fila["campos"][campo] = str(row_cli.get(campo, "") or "").strip()
            fila["destino"] = "consultar"
        elif destino == "corregidos":
            if fila.get("destino") is None and not fila["errores"]:
                fila["destino"] = "corregidos"

    return pending


def _escribir_excel(
    destino: Path | str | BinaryIO,
    filas: list[dict],
    mapeo: dict[str, dict[str, str]],
) -> dict:
    corregidos = [f for f in filas if f.get("destino") == "corregidos"]
    consultar = [f for f in filas if f.get("destino") == "consultar"]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_ok = wb.create_sheet("Corregidos", 0)
    ws_ask = wb.create_sheet("Consultar a Cliente", 1)
    ws_map = wb.create_sheet("Mapeo Potencial Solucion", 2)

    ws_ok.append(CAMPOS)
    _estilo_header(ws_ok, FILL_HEADER_OK)
    for fila in corregidos:
        ws_ok.append([fila["campos"].get(c, "") for c in CAMPOS])
    _estilo_cuerpo(ws_ok)

    cols_map = ["Clave", "RUT ejemplo", "Columna", "Valor antes", "Potencial Solucion"]
    ws_map.append(cols_map)
    _estilo_header(ws_map, FILL_HEADER_MAP)
    mapeo_ordenado = sorted(mapeo.items(), key=lambda x: (x[1]["Columna"], x[1]["Valor antes"]))
    for clave, meta in mapeo_ordenado:
        ws_map.append(
            [
                clave,
                meta["RUT ejemplo"],
                meta["Columna"],
                meta["Valor antes"],
                "",
            ]
        )
    ws_map.column_dimensions["A"].hidden = True
    _estilo_cuerpo(ws_map)

    cols_ask = CAMPOS + ["ERROR", "Potencial Solucion"]
    ws_ask.append(cols_ask)
    _estilo_header(ws_ask, FILL_HEADER_ASK)

    col_index = {name: idx for idx, name in enumerate(cols_ask, start=1)}
    pot_col_letter = get_column_letter(col_index["Potencial Solucion"])

    for fila in consultar:
        error_txt = _mensajes_error(fila.get("errores") or [])
        valores = [fila["campos"].get(c, "") for c in CAMPOS] + [error_txt, ""]
        ws_ask.append(valores)
        excel_row = ws_ask.max_row

        for campo in fila.get("celdas_rojas") or []:
            idx = col_index.get(campo)
            if idx:
                ws_ask.cell(row=excel_row, column=idx).fill = FILL_ERROR_CELL

        claves_info = fila.get("claves_info") or []
        if claves_info:
            if len(claves_info) == 1:
                campo, clave = claves_info[0]
                clave_esc = str(clave).replace('"', '""')
                campo_esc = str(campo).replace('"', '""')
                vlookup = (
                    f'IFERROR(VLOOKUP("{clave_esc}",\'Mapeo Potencial Solucion\'!$A:$E,5,FALSE),"")'
                )
                formula = f'=IF({vlookup}="","","{campo_esc}: "&{vlookup})'
            else:
                chunks = []
                for campo, clave in claves_info:
                    clave_esc = str(clave).replace('"', '""')
                    campo_esc = str(campo).replace('"', '""')
                    vlookup = (
                        f'IFERROR(VLOOKUP("{clave_esc}",\'Mapeo Potencial Solucion\'!$A:$E,5,FALSE),"")'
                    )
                    chunks.append(f'IF({vlookup}="",""," | {campo_esc}: "&{vlookup})')
                formula = f'=MID({("&".join(chunks))},4,9999)'
            ws_ask[f"{pot_col_letter}{excel_row}"] = formula

    _estilo_cuerpo(ws_ask)
    wb.save(destino)
    return {
        "corregidos": len(corregidos),
        "consultar": len(consultar),
        "mapeos": len(mapeo),
    }


def procesar_comparacion(
    ruta_cliente: Path | str,
    ruta_corregida: Path | str,
    destino: Path | str | BinaryIO,
    ruta_analizador: Path | str | None = None,
    decisiones: dict[str, str] | None = None,
) -> dict:
    """
    Procesa comparación (+ analizador opcional).
    Si hay ítems pendientes de decisión y no están resueltos, no escribe Excel y
    retorna status=needs_review con la lista pending.
    """
    filas, mapeo, df_cli, df_corr = _construir_filas(ruta_cliente, ruta_corregida)

    pending: list[dict] = []
    if ruta_analizador:
        items = _leer_analizador(ruta_analizador)
        pending = _aplicar_analizador(filas, df_cli, df_corr, items, decisiones)

    if pending:
        return {
            "status": "needs_review",
            "pending": pending,
            "stats": {
                "cliente": len(df_cli),
                "corregida": len(df_corr),
                "pendientes": len(pending),
            },
        }

    stats = _escribir_excel(destino, filas, mapeo)
    stats["cliente"] = len(df_cli)
    stats["corregida"] = len(df_corr)
    return {"status": "ok", "stats": stats}


def generar_reporte_bytes(
    ruta_cliente: Path | str,
    ruta_corregida: Path | str,
    ruta_analizador: Path | str | None = None,
    decisiones: dict[str, str] | None = None,
) -> dict:
    buf = io.BytesIO()
    result = procesar_comparacion(
        ruta_cliente,
        ruta_corregida,
        buf,
        ruta_analizador=ruta_analizador,
        decisiones=decisiones,
    )
    if result.get("status") == "needs_review":
        return result
    filename = f"Reporte_EMIN_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return {
        "status": "ok",
        "content": buf.getvalue(),
        "filename": filename,
        "stats": result["stats"],
    }
