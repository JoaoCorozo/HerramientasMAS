"""
Análisis de resultados Moodle a partir del Excel de respuestas (informe completo).

Por cada orden N espera el trío:
  Pregunta N | Respuesta N | Respuesta correcta N
(en ES / EN / PT / FR / DE u otras variantes del LMS).

Correcta = texto usuario ≡ respuesta correcta (casefold, sin HTML).
No requiere el archivo de calificaciones.
"""
from __future__ import annotations

import io
import re
import statistics
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

METODOLOGIA = (
    "Se analiza solo el informe de respuestas del LMS. "
    "Cada ítem se identifica por el enunciado (columna Pregunta / Question / …). "
    "Se considera acierto si la respuesta del usuario coincide con la respuesta correcta "
    "(comparación de texto normalizada). "
    "Los nombres de columnas se reconocen en varios idiomas del LMS."
)

# Campos fijos: encabezado normalizado (sin acentos) → clave interna
_HEADER_ALIASES: dict[str, str] = {
    # apellido
    "apellido(s)": "apellido",
    "apellidos": "apellido",
    "apellido": "apellido",
    "surname": "apellido",
    "last name": "apellido",
    "lastname": "apellido",
    "nom de famille": "apellido",
    "nachname": "apellido",
    "sobrenome": "apellido",
    # nombre
    "nombre": "nombre",
    "nombres": "nombre",
    "nome": "nombre",
    "first name": "nombre",
    "firstname": "nombre",
    "prenom": "nombre",
    "vorname": "nombre",
    # RUT / identificador LMS (username suele ser el RUT)
    "rut": "rut",
    "r.u.t.": "rut",
    "r.u.t": "rut",
    "numero de id": "rut",
    "numero de identificacion": "rut",
    "id number": "rut",
    "idnumber": "rut",
    "nombre de usuario": "rut",
    "username": "rut",
    "user name": "rut",
    "identificacao de usuario": "rut",
    "identificacion de usuario": "rut",
    "id de usuario": "rut",
    "login": "rut",
    # email
    "direccion de correo": "email",
    "direccion de correo electronico": "email",
    "correo electronico": "email",
    "correo": "email",
    "email": "email",
    "e mail": "email",
    "email address": "email",
    "adresse de courriel": "email",
    "courriel": "email",
    "e mail adresse": "email",
    "endereco de e mail": "email",
    "endereco de email": "email",
    "endereco eletronic": "email",
    "endereco electronico": "email",

    # estado
    "estado": "estado",
    "state": "estado",
    "status": "estado",
    "etat": "estado",
    "statuszustand": "estado",
    # comenzado
    "comenzado el": "comenzado",
    "comenzado": "comenzado",
    "started on": "comenzado",
    "started": "comenzado",
    "start time": "comenzado",
    "commence le": "comenzado",
    "date de debut": "comenzado",
    "gestartet am": "comenzado",
    "iniciado em": "comenzado",
    "iniciado": "comenzado",
    # finalizado
    "finalizado": "finalizado",
    "finalizado el": "finalizado",
    "completed": "finalizado",
    "completed on": "finalizado",
    "time finished": "finalizado",
    "finished on": "finalizado",
    "termine le": "finalizado",
    "beendet": "finalizado",
    "concluido em": "finalizado",
    "concluido": "finalizado",
    "completo": "finalizado",
    "completa": "finalizado",
    # tiempo
    "tiempo requerido": "tiempo",
    "tiempo empleado": "tiempo",
    "time taken": "tiempo",
    "duration": "tiempo",
    "temps requis": "tiempo",
    "temps utilise": "tiempo",
    "benotigte zeit": "tiempo",
    "tempo utilizado": "tiempo",
    "tempo gasto": "tiempo",
    # comentario
    "comentario": "comentario",
    "comentario -": "comentario",
    "comentarios": "comentario",
    "comment": "comentario",
    "comments": "comentario",
    "commentaire": "comentario",
    "kommentar": "comentario",
    "observacao": "comentario",
    "observacoes": "comentario",
}

# Tríos numerados (orden: correcta antes que respuesta genérica)
_CORRECTA_PATTERNS = [
    re.compile(r"^respuesta\s+correcta\s*(\d+)$"),
    re.compile(r"^resposta\s+correta\s*(\d+)$"),
    re.compile(r"^right\s+answer\s*(\d+)$"),
    re.compile(r"^correct\s+answer\s*(\d+)$"),
    re.compile(r"^correct\s+response\s*(\d+)$"),
    re.compile(r"^bonne\s+reponse\s*(\d+)$"),
    re.compile(r"^reponse\s+correcte\s*(\d+)$"),
    re.compile(r"^richtige\s+antwort\s*(\d+)$"),
    re.compile(r"^antwort\s+richtig\s*(\d+)$"),
]
_PREGUNTA_PATTERNS = [
    re.compile(r"^pregunta\s*(\d+)$"),
    re.compile(r"^question\s*(\d+)$"),
    re.compile(r"^pergunta\s*(\d+)$"),
    re.compile(r"^questao\s*(\d+)$"),  # PT-BR Moodle: Questão N
    re.compile(r"^frage\s*(\d+)$"),
]
_RESPUESTA_PATTERNS = [
    re.compile(r"^respuesta\s*(\d+)$"),
    re.compile(r"^response\s*(\d+)$"),
    re.compile(r"^resposta\s*(\d+)$"),
    re.compile(r"^answer\s*(\d+)$"),
    re.compile(r"^reponse\s*(\d+)$"),
    re.compile(r"^antwort\s*(\d+)$"),
]
_NOTA_RE = re.compile(
    r"^(calificacion|classificacao|grade|nota|mark|note|bewertung|pontuacao|avaliar|avaliacao)\s*/\s*[\d.,]+$",
)
_HTML_RE = re.compile(r"<[^>]+>")
_TIME_RE = re.compile(
    r"(\d+)\s*"
    r"(horas?|hours?|heures?|stunden?|"
    r"minutos?|minutes?|minuten?|"
    r"segundos?|seconds?|secondes?|sekunden?)",
    re.I,
)


def _strip_accents(s: str) -> str:
    nfkd = unicodedata.normalize("NFKD", s)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def _norm_header(val: Any) -> str:
    s = str(val or "").strip().lower().replace("\xa0", " ")
    s = _strip_accents(s)
    s = s.replace("_", " ").replace("-", " ")
    s = re.sub(r"\s+", " ", s).strip()
    # "Comentario -" / "Comment -" residual
    s = re.sub(r"\s+$", "", s)
    if s.endswith(" -"):
        s = s[:-2].strip()
    return s


def _clean_text(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    s = _HTML_RE.sub(" ", s)
    s = s.replace("&nbsp;", " ").replace("&amp;", "&")
    return re.sub(r"\s+", " ", s).strip()


def _parse_num(val: Any) -> float | None:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _textos_equivalentes(a: str, b: str) -> bool:
    return _clean_text(a).casefold() == _clean_text(b).casefold()


def _parse_tiempo_segundos(texto: str) -> int | None:
    if not texto:
        return None
    total = 0
    found = False
    for num, unit in _TIME_RE.findall(texto):
        found = True
        n = int(num)
        u = _strip_accents(unit.lower())
        if u.startswith(("hora", "hour", "heure", "stund")):
            total += n * 3600
        elif u.startswith(("minuto", "minute", "minut")):
            total += n * 60
        else:
            total += n
    return total if found else None


def _fmt_tiempo(seg: int | None) -> str:
    if seg is None:
        return ""
    h, rem = divmod(seg, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h}h {m}m {s}s"
    if m:
        return f"{m}m {s}s"
    return f"{s}s"


@dataclass
class ItemRespuesta:
    orden: int
    enunciado: str
    respuesta_usuario: str
    respuesta_correcta: str
    correcta: bool | None  # None si falta enunciado o respuesta correcta


@dataclass
class Intento:
    apellido: str
    nombre: str
    email: str
    rut: str
    estado: str
    comenzado: str
    finalizado: str
    tiempo: str
    tiempo_seg: int | None
    nota: float | None
    comentario: str
    items: list[ItemRespuesta] = field(default_factory=list)
    anon_id: str = ""

    @property
    def clave(self) -> tuple[str, str]:
        return (self.email.lower().strip(), self.comenzado.strip())

    @property
    def nombre_completo(self) -> str:
        return f"{self.nombre} {self.apellido}".strip()

    @property
    def correctas(self) -> int:
        return sum(1 for x in self.items if x.correcta is True)

    @property
    def incorrectas(self) -> int:
        return sum(1 for x in self.items if x.correcta is False)


def _map_headers(header_row: tuple[Any, ...]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        h = _norm_header(cell)
        if not h:
            continue
        if h in _HEADER_ALIASES:
            mapping[_HEADER_ALIASES[h]] = i
            continue
        if _NOTA_RE.match(h) or h in ("calificacion", "classificacao", "grade", "nota", "mark", "note", "bewertung"):
            mapping["nota"] = i
            continue
        # Fallback flexible: correo / email en cualquier idioma
        if "email" not in mapping and any(
            tok in h for tok in ("mail", "correo", "courriel", "e mail")
        ):
            mapping["email"] = i
            continue
        # Fallback: username / identificación de usuario → RUT
        if "rut" not in mapping and (
            h in ("usuario", "user")
            or "nombre de usuario" in h
            or "identificacao de usuario" in h
            or "identificacion de usuario" in h
            or h.endswith(" de usuario")
            or h == "username"
        ):
            mapping["rut"] = i
            continue
        matched = False
        for pat in _CORRECTA_PATTERNS:
            m = pat.match(h)
            if m:
                mapping[f"correcta_{int(m.group(1))}"] = i
                matched = True
                break
        if matched:
            continue
        for pat in _PREGUNTA_PATTERNS:
            m = pat.match(h)
            if m:
                mapping[f"pregunta_{int(m.group(1))}"] = i
                matched = True
                break
        if matched:
            continue
        for pat in _RESPUESTA_PATTERNS:
            m = pat.match(h)
            if m:
                mapping[f"respuesta_{int(m.group(1))}"] = i
                break
    return mapping


def _leer_hoja(ruta: Path | str) -> list[tuple[Any, ...]]:
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        return [tuple(row) for row in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


def _filas_utiles(rows: list[tuple[Any, ...]]) -> tuple[dict[str, int], list[tuple[Any, ...]], list[str]]:
    if not rows:
        raise ValueError("El Excel está vacío.")
    mapping = _map_headers(rows[0])
    headers_raw = [_norm_header(c) for c in rows[0] if c is not None and str(c).strip()]
    if "email" not in mapping:
        raise ValueError(
            "No se encontró la columna de correo (Dirección de correo / Email address / …). "
            f"Encabezados detectados: {', '.join(headers_raw[:12])}…"
        )
    data = []
    for row in rows[1:]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        email = _clean_text(row[mapping["email"]] if mapping["email"] < len(row) else "")
        if not email or "@" not in email:
            continue
        data.append(row)
    return mapping, data, headers_raw


def _get(row: tuple[Any, ...], mapping: dict[str, int], key: str) -> Any:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _detectar_ordenes(map_r: dict[str, int]) -> list[int]:
    nums: set[int] = set()
    for k in map_r:
        for pref in ("pregunta_", "respuesta_", "correcta_"):
            if k.startswith(pref):
                nums.add(int(k.split("_")[1]))
    return sorted(nums)


def analizar_respuestas(ruta_respuestas: Path | str) -> dict[str, Any]:
    map_r, data_r, headers_raw = _filas_utiles(_leer_hoja(ruta_respuestas))
    ordenes = _detectar_ordenes(map_r)
    if not ordenes:
        raise ValueError(
            "No se encontraron columnas de preguntas/respuestas numeradas "
            "(Pregunta N / Question N, Respuesta N / Response N, Respuesta correcta N / Right answer N)."
        )

    incompletos = [
        n
        for n in ordenes
        if not (
            f"pregunta_{n}" in map_r
            and f"respuesta_{n}" in map_r
            and f"correcta_{n}" in map_r
        )
    ]
    ordenes_completos = [n for n in ordenes if n not in incompletos]
    if not ordenes_completos:
        raise ValueError(
            "El Excel debe incluir el trío completo por cada ítem: "
            "Pregunta/Question/Questão N + Respuesta/Response/Resposta N + "
            "Respuesta correcta/Right answer/Resposta correta N."
        )

    # Agrupar intentos; si hay duplicado de (email, comenzado) se conserva el último y se registra
    by_key: dict[tuple[str, str], tuple[Any, ...]] = {}
    dup_keys: list[str] = []
    for row in data_r:
        email = _clean_text(_get(row, map_r, "email"))
        comenzado = _clean_text(_get(row, map_r, "comenzado"))
        key = (email.lower(), comenzado)
        if key in by_key:
            dup_keys.append(f"{email} | {comenzado}")
        by_key[key] = row

    intentos: list[Intento] = []
    items_sin_evaluar = 0
    for idx, (key, row) in enumerate(sorted(by_key.items()), start=1):
        tiempo = _clean_text(_get(row, map_r, "tiempo"))
        items: list[ItemRespuesta] = []
        for n in ordenes_completos:
            enunciado = _clean_text(_get(row, map_r, f"pregunta_{n}"))
            resp_u = _clean_text(_get(row, map_r, f"respuesta_{n}"))
            resp_ok = _clean_text(_get(row, map_r, f"correcta_{n}"))
            if not enunciado or not resp_ok:
                items_sin_evaluar += 1
                correcta = None
            else:
                correcta = _textos_equivalentes(resp_u, resp_ok)
            items.append(
                ItemRespuesta(
                    orden=n,
                    enunciado=enunciado,
                    respuesta_usuario=resp_u,
                    respuesta_correcta=resp_ok,
                    correcta=correcta,
                )
            )

        intentos.append(
            Intento(
                apellido=_clean_text(_get(row, map_r, "apellido")),
                nombre=_clean_text(_get(row, map_r, "nombre")),
                email=_clean_text(_get(row, map_r, "email")),
                rut=_clean_text(_get(row, map_r, "rut")),
                estado=_clean_text(_get(row, map_r, "estado")),
                comenzado=_clean_text(_get(row, map_r, "comenzado")),
                finalizado=_clean_text(_get(row, map_r, "finalizado")),
                tiempo=tiempo,
                tiempo_seg=_parse_tiempo_segundos(tiempo),
                nota=_parse_num(_get(row, map_r, "nota")),
                comentario=_clean_text(_get(row, map_r, "comentario")),
                items=items,
                anon_id=f"Usuario {idx:03d}",
            )
        )

    intentos_nom = sorted(intentos, key=lambda i: (i.apellido.lower(), i.nombre.lower(), i.email))

    notas = [i.nota for i in intentos if i.nota is not None]
    tiempos = [i.tiempo_seg for i in intentos if i.tiempo_seg is not None]
    aprobados = sum(1 for n in notas if n >= 60)
    total_eval = sum(1 for i in intentos for x in i.items if x.correcta is not None)
    total_ok = sum(i.correctas for i in intentos)
    total_bad = sum(i.incorrectas for i in intentos)

    por_pregunta: dict[str, dict[str, Any]] = {}
    for it in intentos:
        for item in it.items:
            if not item.enunciado:
                continue
            st = por_pregunta.setdefault(
                item.enunciado,
                {
                    "enunciado": item.enunciado,
                    "respuesta_correcta": item.respuesta_correcta,
                    "veces_presentada": 0,
                    "aciertos": 0,
                    "errores": 0,
                    "sin_evaluar": 0,
                    "respuestas_usuario": Counter(),
                    "respuestas_incorrectas": Counter(),
                    "personas": set(),
                    "ordenes": set(),
                },
            )
            if item.respuesta_correcta and not st["respuesta_correcta"]:
                st["respuesta_correcta"] = item.respuesta_correcta
            st["veces_presentada"] += 1
            st["personas"].add(it.clave)
            st["ordenes"].add(item.orden)
            if item.respuesta_usuario:
                st["respuestas_usuario"][item.respuesta_usuario] += 1
            if item.correcta is True:
                st["aciertos"] += 1
            elif item.correcta is False:
                st["errores"] += 1
                if item.respuesta_usuario:
                    st["respuestas_incorrectas"][item.respuesta_usuario] += 1
            else:
                st["sin_evaluar"] += 1

    preguntas: list[dict[str, Any]] = []
    for enunciado, st in por_pregunta.items():
        correctas_vistas: Counter[str] = Counter()
        for it in intentos:
            for item in it.items:
                if item.enunciado == enunciado and item.respuesta_correcta:
                    correctas_vistas[item.respuesta_correcta] += 1
        resp_ok = correctas_vistas.most_common(1)[0][0] if correctas_vistas else st["respuesta_correcta"]
        evaluadas = st["aciertos"] + st["errores"]
        # Todas las incorrectas (ordenadas por frecuencia) para expandir en Excel
        top_malas = [{"respuesta": t, "veces": n} for t, n in st["respuestas_incorrectas"].most_common()]
        top_todas = [{"respuesta": t, "veces": n} for t, n in st["respuestas_usuario"].most_common(10)]
        preguntas.append(
            {
                "enunciado": enunciado,
                "respuesta_correcta": resp_ok,
                "veces_presentada": st["veces_presentada"],
                "personas_distintas": len(st["personas"]),
                "aciertos": st["aciertos"],
                "errores": st["errores"],
                "pct_acierto": round(100.0 * st["aciertos"] / evaluadas, 1) if evaluadas else None,
                "aparecio_en_ordenes_P": sorted(st["ordenes"]),
                "distractores_mas_elegidos": top_malas,
                "respuestas_mas_frecuentes": top_todas,
                "incoherencias_cruce": 0,
            }
        )

    preguntas.sort(
        key=lambda x: (x["pct_acierto"] if x["pct_acierto"] is not None else 999, -x["errores"], x["enunciado"])
    )
    preguntas_mas_dificiles = [p for p in preguntas if p["errores"] > 0][:25]

    # Detalle largo: 1 fila por usuario × pregunta (filtro/pivote en Excel)
    detalle_respuestas: list[dict[str, Any]] = []
    for it in intentos_nom:
        for item in it.items:
            detalle_respuestas.append(
                {
                    "rut": it.rut,
                    "apellido": it.apellido,
                    "nombre": it.nombre,
                    "email": it.email,
                    "comenzado": it.comenzado,
                    "finalizado": it.finalizado,
                    "nota": it.nota,
                    "tiempo": it.tiempo,
                    "orden": item.orden,
                    "enunciado": item.enunciado,
                    "respuesta_usuario": item.respuesta_usuario,
                    "respuesta_correcta": item.respuesta_correcta,
                    "resultado": (
                        "Correcta"
                        if item.correcta is True
                        else ("Incorrecta" if item.correcta is False else "Sin evaluar")
                    ),
                }
            )

    def _row_intento(it: Intento) -> dict[str, Any]:
        return {
            "rut": it.rut,
            "apellido": it.apellido,
            "nombre": it.nombre,
            "nombre_completo": it.nombre_completo,
            "email": it.email,
            "nota": it.nota,
            "correctas": it.correctas,
            "incorrectas": it.incorrectas,
            "total_preguntas_intento": len(it.items),
            "pct_acierto_intento": round(100.0 * it.correctas / len(it.items), 1) if it.items else None,
            "tiempo": it.tiempo,
            "estado": it.estado,
            "comenzado": it.comenzado,
            "finalizado": it.finalizado,
        }
    columnas_detectadas = {
        "email": "email" in map_r,
        "rut": "rut" in map_r,
        "comenzado": "comenzado" in map_r,
        "nota": "nota" in map_r,
        "ordenes_completos": ordenes_completos,
        "ordenes_incompletos": incompletos,
        "ejemplos_encabezados": headers_raw[:30],
    }

    resumen = {
        "intentos_emparejados": len(intentos),
        "filas_respuestas": len(data_r),
        "filas_calificaciones": 0,
        "preguntas_por_intento": len(ordenes_completos),
        "preguntas_distintas_banco": len(preguntas),
        "respuestas_evaluadas": total_eval,
        "aciertos_totales": total_ok,
        "errores_totales": total_bad,
        "pct_acierto_global": round(100.0 * total_ok / total_eval, 1) if total_eval else None,
        "incoherencias_texto_vs_puntaje": 0,
        "items_sin_evaluar": items_sin_evaluar,
        "nota_promedio": round(statistics.mean(notas), 2) if notas else None,
        "nota_mediana": round(statistics.median(notas), 2) if notas else None,
        "nota_min": min(notas) if notas else None,
        "nota_max": max(notas) if notas else None,
        "aprobados_ge_60": aprobados,
        "pct_aprobacion_ge_60": round(100.0 * aprobados / len(notas), 1) if notas else None,
        "tiempo_promedio": _fmt_tiempo(int(statistics.mean(tiempos))) if tiempos else None,
        "tiempo_mediana": _fmt_tiempo(int(statistics.median(tiempos))) if tiempos else None,
        "tiempo_min": _fmt_tiempo(min(tiempos)) if tiempos else None,
        "tiempo_max": _fmt_tiempo(max(tiempos)) if tiempos else None,
        "metodologia": METODOLOGIA,
        "regla_puntaje": "Acierto = respuesta del usuario igual a respuesta correcta (texto)",
        "columnas_detectadas": columnas_detectadas,
    }

    validacion = {
        "ok": not dup_keys and not incompletos and items_sin_evaluar == 0,
        "solo_en_respuestas": [],
        "solo_en_calificaciones": [],
        "duplicados_respuestas": dup_keys[:50],
        "duplicados_calificaciones": [],
        "discrepancias_nota": [],
        "incoherencias_texto_vs_puntaje": [],
        "ordenes_incompletos": incompletos,
        "items_sin_evaluar": items_sin_evaluar,
        "total_solo_respuestas": 0,
        "total_solo_calificaciones": 0,
        "total_incoherencias": 0,
        "total_duplicados": len(dup_keys),
    }

    return {
        "resumen": resumen,
        "validacion": validacion,
        "metodologia": METODOLOGIA,
        "preguntas": preguntas,
        "preguntas_mas_dificiles": preguntas_mas_dificiles,
        "intentos_nominativos": [_row_intento(it) for it in intentos_nom],
        "detalle_respuestas": detalle_respuestas,
        "preview_nominativos": [_row_intento(it) for it in intentos_nom[:20]],
        "preview_preguntas_dificiles": preguntas_mas_dificiles[:12],
        "preview_detalle": detalle_respuestas[:30],
    }


# Compatibilidad con import antiguo
def analizar_archivos(ruta_respuestas: Path | str, ruta_calificaciones: Path | str | None = None) -> dict[str, Any]:
    """Solo usa respuestas; ruta_calificaciones se ignora si se pasa."""
    return analizar_respuestas(ruta_respuestas)


def _style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, max_width: int = 56) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max((len(str(c.value or "")) for c in col[:80]), default=8)
        ws.column_dimensions[letter].width = min(max(10, length + 2), max_width)


def _write_table(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_header(ws)
    _autosize(ws)


def build_excel_analisis(resultado: dict[str, Any]) -> tuple[bytes, str]:
    """Excel lean: Resumen | Por_pregunta | Por_usuario | Detalle_respuestas | Validacion."""
    wb = openpyxl.Workbook()
    r = resultado["resumen"]
    v = resultado["validacion"]

    ws = wb.active
    ws.title = "Resumen"
    _write_table(
        ws,
        ["Métrica", "Valor"],
        [
            ["Intentos analizados", r["intentos_emparejados"]],
            ["Preguntas distintas del banco", r["preguntas_distintas_banco"]],
            ["Preguntas por intento", r["preguntas_por_intento"]],
            ["Respuestas evaluadas", r["respuestas_evaluadas"]],
            ["Aciertos / errores", f"{r['aciertos_totales']} / {r['errores_totales']}"],
            ["% acierto global", r["pct_acierto_global"]],
            ["Nota promedio / mediana", f"{r['nota_promedio']} / {r['nota_mediana']}"],
            ["Nota mín / máx", f"{r['nota_min']} / {r['nota_max']}"],
            ["% aprobación ≥60", r["pct_aprobacion_ge_60"]],
            ["Tiempo promedio", r["tiempo_promedio"]],
            ["Criterio acierto", r["regla_puntaje"]],
            ["Filas en Detalle_respuestas", len(resultado.get("detalle_respuestas") or [])],
            ["Validación OK", "Sí" if v["ok"] else "No"],
            ["Metodología", r["metodologia"]],
        ],
    )

    # Ordenado de más difícil a más fácil.
    # Por cada respuesta incorrecta distinta: columna texto + columna cantidad.
    ws_p = wb.create_sheet("Por_pregunta")
    preguntas = resultado["preguntas"]
    max_incorrectas = max(
        (len(p.get("distractores_mas_elegidos") or []) for p in preguntas),
        default=0,
    )
    headers_p = [
        "Enunciado",
        "Respuesta correcta",
        "Veces presentada",
        "Aciertos",
        "Errores",
        "% acierto",
    ]
    for i in range(1, max_incorrectas + 1):
        headers_p.append(f"Respuesta incorrecta {i}")
        headers_p.append(f"Veces {i}")
    rows_p: list[list[Any]] = []
    for p in preguntas:
        row: list[Any] = [
            p["enunciado"],
            p["respuesta_correcta"],
            p["veces_presentada"],
            p["aciertos"],
            p["errores"],
            p["pct_acierto"],
        ]
        distractores = p.get("distractores_mas_elegidos") or []
        for i in range(max_incorrectas):
            if i < len(distractores):
                row.append(distractores[i]["respuesta"])
                row.append(distractores[i]["veces"])
            else:
                row.append("")
                row.append("")
        rows_p.append(row)
    _write_table(ws_p, headers_p, rows_p)

    ws_u = wb.create_sheet("Por_usuario")
    _write_table(
        ws_u,
        [
            "RUT",
            "Apellido",
            "Nombre",
            "Email",
            "Nota",
            "Correctas",
            "Incorrectas",
            "% acierto",
            "Tiempo",
            "Comenzado",
        ],
        [
            [
                it.get("rut") or "",
                it["apellido"],
                it["nombre"],
                it["email"],
                it["nota"],
                it["correctas"],
                it["incorrectas"],
                it["pct_acierto_intento"],
                it["tiempo"],
                it["comenzado"],
            ]
            for it in resultado["intentos_nominativos"]
        ],
    )

    # Hoja principal de auditoría: filtrar por RUT / email / enunciado / Incorrecta
    ws_d = wb.create_sheet("Detalle_respuestas")
    detalle = resultado.get("detalle_respuestas") or []
    _write_table(
        ws_d,
        [
            "RUT",
            "Apellido",
            "Nombre",
            "Email",
            "Comenzado",
            "Nota intento",
            "Orden",
            "Enunciado",
            "Respuesta usuario",
            "Respuesta correcta",
            "Resultado",
        ],
        [
            [
                d.get("rut") or "",
                d["apellido"],
                d["nombre"],
                d["email"],
                d["comenzado"],
                d["nota"],
                d["orden"],
                d["enunciado"],
                d["respuesta_usuario"],
                d["respuesta_correcta"],
                d["resultado"],
            ]
            for d in detalle
        ],
    )

    if not v.get("ok"):
        ws_v = wb.create_sheet("Validacion")
        ws_v.append(["Tipo", "Detalle"])
        ws_v.append(["Duplicados email+comenzado", str(v.get("total_duplicados", 0))])
        ws_v.append(
            [
                "Órdenes incompletos",
                ", ".join(map(str, v.get("ordenes_incompletos") or [])) or "(ninguno)",
            ]
        )
        ws_v.append(["Ítems sin evaluar", v.get("items_sin_evaluar", 0)])
        for item in v.get("duplicados_respuestas") or []:
            ws_v.append(["Duplicado", item])
        _style_header(ws_v)
        _autosize(ws_v)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "Analisis_Resultados_Evaluacion.xlsx"
